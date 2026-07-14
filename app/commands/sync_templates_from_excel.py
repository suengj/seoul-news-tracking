"""Synchronize `config/message_templates.yaml` from the Excel workbook.

Excel (`templates/서울시_재난특보_X템플릿.xlsx`) is the human-managed business
source. The generated YAML is the runtime snapshot loaded by Telegram,
renderer, Rule, and AI paths — never open the workbook on every request.

Usage:
    python -m app.commands.sync_templates_from_excel --check
    python -m app.commands.sync_templates_from_excel --write
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import load_workbook

from app.config import PROJECT_ROOT, TEMPLATES_PATH
from app.template_ids import (
    EXPECTED_AUTOMATION_COUNT,
    EXPECTED_REFERENCE_COUNT,
    EXPECTED_WORKBOOK_TOTAL,
    LEGACY_ALIAS_TO_CANONICAL,
)
from app.template_renderer import clear_template_cache, load_templates

DEFAULT_WORKBOOK = PROJECT_ROOT / "templates" / "서울시_재난특보_X템플릿.xlsx"
CATALOG_SHEET = "01_템플릿목록"
WORDING_SHEET = "02_문안템플릿"

CATALOG_COLUMNS = ("템플릿ID", "재해유형", "등급", "상태", "발동 조건", "검수상태", "자동화")
WORDING_COLUMNS = ("템플릿ID", "구분", "검수상태", "게시 문안 (변수는 {중괄호})", "사용 변수")

# Business placeholders only — ignore renderer control markers like `{#slot}`.
_PLACEHOLDER_RE = re.compile(r"\{([^{}#/][^{}]*)\}")
# Compact button labels where level+state alone is too long / ambiguous.
BUTTON_LABEL_OVERRIDES: dict[str, str] = {
    "HW-05": "특보 해제",
    "HW-06": "예비특보",
    "HW-07": "교통 연계",
    "HW-08": "홍수 복합",
    "HW-09": "태풍 연계",
    "HT-05": "특보 해제",
    "TN-02": "주의보 해제",
    "FL-03": "특보 해제",
}

CANONICAL_TO_ALIASES: dict[str, list[str]] = {}
for _legacy, _canonical in LEGACY_ALIAS_TO_CANONICAL.items():
    CANONICAL_TO_ALIASES.setdefault(_canonical, []).append(_legacy)


@dataclass
class ValidationIssue:
    template_id: str
    sheet: str
    row: int
    field: str
    reason: str

    def format(self) -> str:
        tid = self.template_id or "(blank)"
        return (
            f"[{self.sheet} row {self.row}] id={tid} field={self.field}: {self.reason}"
        )


@dataclass
class CatalogRow:
    template_id: str
    disaster_type: str
    level: str
    state: str
    trigger_condition: str
    review_label: str
    automation_raw: str
    row: int


@dataclass
class WordingRow:
    template_id: str
    section: str
    review_label: str
    text: str
    variables_raw: str
    row: int


@dataclass
class SyncResult:
    templates: list[dict] = field(default_factory=list)
    issues: list[ValidationIssue] = field(default_factory=list)
    total: int = 0
    automation_count: int = 0
    reference_count: int = 0

    @property
    def ok(self) -> bool:
        return not self.issues


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() if isinstance(value, str) else str(value).strip()


def _header_index(ws, expected: tuple[str, ...]) -> dict[str, int]:
    headers = {_cell_str(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}
    missing = [name for name in expected if name not in headers]
    if missing:
        raise ValueError(f"missing required column(s): {', '.join(missing)}")
    return headers


def _read_catalog(ws) -> tuple[list[CatalogRow], list[ValidationIssue]]:
    headers = _header_index(ws, CATALOG_COLUMNS)
    rows: list[CatalogRow] = []
    issues: list[ValidationIssue] = []
    seen: set[str] = set()
    for r in range(2, ws.max_row + 1):
        tid = _cell_str(ws.cell(r, headers["템플릿ID"]).value)
        if not any(
            _cell_str(ws.cell(r, headers[col]).value) for col in CATALOG_COLUMNS
        ):
            continue
        if not tid:
            issues.append(
                ValidationIssue("", CATALOG_SHEET, r, "템플릿ID", "blank template ID")
            )
            continue
        if tid in seen:
            issues.append(
                ValidationIssue(tid, CATALOG_SHEET, r, "템플릿ID", "duplicate template ID")
            )
            continue
        seen.add(tid)
        automation = _cell_str(ws.cell(r, headers["자동화"]).value)
        if automation not in ("대상", "제외"):
            issues.append(
                ValidationIssue(
                    tid, CATALOG_SHEET, r, "자동화", f"must be 대상 or 제외, got {automation!r}"
                )
            )
        rows.append(
            CatalogRow(
                template_id=tid,
                disaster_type=_cell_str(ws.cell(r, headers["재해유형"]).value),
                level=_cell_str(ws.cell(r, headers["등급"]).value),
                state=_cell_str(ws.cell(r, headers["상태"]).value),
                trigger_condition=_cell_str(ws.cell(r, headers["발동 조건"]).value),
                review_label=_cell_str(ws.cell(r, headers["검수상태"]).value),
                automation_raw=automation,
                row=r,
            )
        )
    return rows, issues


def _read_wording(ws) -> tuple[list[WordingRow], list[ValidationIssue]]:
    headers = _header_index(ws, WORDING_COLUMNS)
    rows: list[WordingRow] = []
    issues: list[ValidationIssue] = []
    seen: set[str] = set()
    text_col = "게시 문안 (변수는 {중괄호})"
    for r in range(2, ws.max_row + 1):
        tid = _cell_str(ws.cell(r, headers["템플릿ID"]).value)
        if not any(_cell_str(ws.cell(r, headers[col]).value) for col in WORDING_COLUMNS):
            continue
        if not tid:
            issues.append(
                ValidationIssue("", WORDING_SHEET, r, "템플릿ID", "blank template ID")
            )
            continue
        if tid in seen:
            issues.append(
                ValidationIssue(tid, WORDING_SHEET, r, "템플릿ID", "duplicate template ID")
            )
            continue
        seen.add(tid)
        text = ws.cell(r, headers[text_col]).value
        text_str = "" if text is None else str(text)
        # Preserve workbook newlines; strip only a single trailing newline if
        # Excel stored one — YAML `|` will re-terminate lines consistently.
        if text_str.endswith("\r\n"):
            text_str = text_str[:-2]
        elif text_str.endswith("\n"):
            text_str = text_str[:-1]
        rows.append(
            WordingRow(
                template_id=tid,
                section=_cell_str(ws.cell(r, headers["구분"]).value),
                review_label=_cell_str(ws.cell(r, headers["검수상태"]).value),
                text=text_str,
                variables_raw=_cell_str(ws.cell(r, headers["사용 변수"]).value),
                row=r,
            )
        )
    return rows, issues


def extract_placeholders(text: str) -> list[str]:
    """Conservative `{변수명}` parser — skips `{#slot}` / `{/slot}` control markers."""
    return _PLACEHOLDER_RE.findall(text)


def parse_declared_variables(raw: str) -> list[str]:
    if not raw:
        return []
    found = extract_placeholders(raw)
    if found:
        return found
    # Fallback: space-separated bare names (should not occur in current workbook).
    return [part.strip("{} ") for part in raw.split() if part.strip()]


def normalize_review_status(label: str) -> str:
    if any(token in label for token in ("확정", "변경금지", "게시이력")):
        return "confirmed"
    if any(token in label for token in ("검수필요", "부서검수", "★")):
        return "review_required"
    if "참고" in label:
        return "reference"
    return "review_required" if label else "confirmed"


def build_display_name(disaster: str, level: str, state: str) -> str:
    if not level or level == "-":
        if state == "해제":
            return f"{disaster}특보 해제"
        return f"{disaster} {state}".strip()
    if "/" in level:
        return f"{disaster}특보 {state}".strip()
    return f"{disaster}{level} {state}".strip()


def build_button_label(template_id: str, level: str, state: str) -> str:
    if template_id in BUTTON_LABEL_OVERRIDES:
        return BUTTON_LABEL_OVERRIDES[template_id]
    if level and level != "-":
        short_level = level.split("/")[0]
        return f"{short_level} {state}".strip()
    return state


def category_code(template_id: str) -> str:
    prefix = template_id.split("-", 1)[0]
    return prefix


def _validate_wording_vs_vars(
    wording: WordingRow, automation: bool, issues: list[ValidationIssue]
) -> list[str]:
    text = wording.text
    if automation and not text.strip():
        issues.append(
            ValidationIssue(
                wording.template_id,
                WORDING_SHEET,
                wording.row,
                "게시 문안",
                "automation template has empty message",
            )
        )
    # Unbalanced / truncated braces (simple check for bare leftovers).
    if text.count("{") != text.count("}"):
        issues.append(
            ValidationIssue(
                wording.template_id,
                WORDING_SHEET,
                wording.row,
                "게시 문안",
                "unbalanced braces in message",
            )
        )
    placeholders = extract_placeholders(text)
    declared = parse_declared_variables(wording.variables_raw)
    ph_set, decl_set = set(placeholders), set(declared)
    for name in sorted(ph_set - decl_set):
        issues.append(
            ValidationIssue(
                wording.template_id,
                WORDING_SHEET,
                wording.row,
                "게시 문안",
                f"placeholder {{{name}}} not listed in 사용 변수",
            )
        )
    for name in sorted(decl_set - ph_set):
        issues.append(
            ValidationIssue(
                wording.template_id,
                WORDING_SHEET,
                wording.row,
                "사용 변수",
                f"variable {{{name}}} absent from 게시 문안",
            )
        )
    # Prefer 사용 변수 order when it matches; otherwise text order.
    ordered = [n for n in declared if n in ph_set] or placeholders
    return ordered


def build_from_workbook(workbook_path: Path) -> SyncResult:
    result = SyncResult()
    try:
        wb = load_workbook(workbook_path, data_only=True)
    except Exception as exc:  # noqa: BLE001 — surface as validation failure
        result.issues.append(
            ValidationIssue("", str(workbook_path), 0, "workbook", f"cannot open: {exc}")
        )
        return result

    if CATALOG_SHEET not in wb.sheetnames:
        result.issues.append(
            ValidationIssue("", "workbook", 0, "sheet", f"missing sheet {CATALOG_SHEET!r}")
        )
    if WORDING_SHEET not in wb.sheetnames:
        result.issues.append(
            ValidationIssue("", "workbook", 0, "sheet", f"missing sheet {WORDING_SHEET!r}")
        )
    if result.issues:
        return result

    try:
        catalog_rows, catalog_issues = _read_catalog(wb[CATALOG_SHEET])
        wording_rows, wording_issues = _read_wording(wb[WORDING_SHEET])
    except ValueError as exc:
        result.issues.append(
            ValidationIssue("", "workbook", 1, "columns", str(exc))
        )
        return result

    result.issues.extend(catalog_issues)
    result.issues.extend(wording_issues)

    catalog_by_id = {row.template_id: row for row in catalog_rows}
    wording_by_id = {row.template_id: row for row in wording_rows}

    for tid, crow in catalog_by_id.items():
        if tid not in wording_by_id:
            result.issues.append(
                ValidationIssue(
                    tid, CATALOG_SHEET, crow.row, "템플릿ID", "catalog ID has no wording row"
                )
            )
    for tid, wrow in wording_by_id.items():
        if tid not in catalog_by_id:
            result.issues.append(
                ValidationIssue(
                    tid, WORDING_SHEET, wrow.row, "템플릿ID", "wording row has no catalog row"
                )
            )

    joined: list[dict] = []
    for crow in catalog_rows:
        wrow = wording_by_id.get(crow.template_id)
        if wrow is None:
            continue
        automation = crow.automation_raw == "대상"
        slots = _validate_wording_vs_vars(wrow, automation, result.issues)
        # Prefer wording-sheet review label (more specific) when present.
        review_label = wrow.review_label or crow.review_label
        entry = {
            "id": crow.template_id,
            "category_code": category_code(crow.template_id),
            "disaster_type": crow.disaster_type,
            "level": crow.level,
            "state": crow.state,
            "display_name": build_display_name(crow.disaster_type, crow.level, crow.state),
            "button_label": build_button_label(crow.template_id, crow.level, crow.state),
            "section": wrow.section,
            "trigger_condition": crow.trigger_condition,
            "review_status": normalize_review_status(review_label),
            "review_label": review_label,
            "automation": automation,
            "enabled": automation,
            "menu_order": crow.row,
            "aliases": list(CANONICAL_TO_ALIASES.get(crow.template_id, [])),
            "required_slots": slots,
            "optional_slots": [],
            "text": wrow.text,
        }
        joined.append(entry)

    result.templates = joined
    result.total = len(joined)
    result.automation_count = sum(1 for t in joined if t["automation"])
    result.reference_count = sum(1 for t in joined if not t["automation"])

    if result.total != EXPECTED_WORKBOOK_TOTAL:
        result.issues.append(
            ValidationIssue(
                "",
                "workbook",
                0,
                "count",
                f"expected {EXPECTED_WORKBOOK_TOTAL} templates, got {result.total}",
            )
        )
    if result.automation_count != EXPECTED_AUTOMATION_COUNT:
        result.issues.append(
            ValidationIssue(
                "",
                "workbook",
                0,
                "count",
                f"expected {EXPECTED_AUTOMATION_COUNT} automation templates, "
                f"got {result.automation_count}",
            )
        )
    if result.reference_count != EXPECTED_REFERENCE_COUNT:
        result.issues.append(
            ValidationIssue(
                "",
                "workbook",
                0,
                "count",
                f"expected {EXPECTED_REFERENCE_COUNT} reference templates, "
                f"got {result.reference_count}",
            )
        )
    return result


def _system_and_legacy_sections() -> dict:
    """Preserve service-control + one legacy-hidden template beside Excel catalog."""
    return {
        "system_templates": [
            {
                "id": "ORIGINAL_ONLY",
                "display_name": "원문 그대로",
                "button_label": "📄 원문",
                "enabled": True,
                "automation": False,
                "review_status": "confirmed",
                "review_label": "system",
                "aliases": [],
                "required_slots": ["원문"],
                "optional_slots": [],
                "text": "{원문}\n",
            },
            {
                "id": "UNKNOWN",
                "display_name": "권장 템플릿 없음",
                "button_label": "",
                "enabled": False,
                "automation": False,
                "review_status": "confirmed",
                "review_label": "system",
                "aliases": [],
                "required_slots": [],
                "optional_slots": [],
                "text": "",
            },
        ],
        "legacy_templates": [
            {
                "id": "HEAVY_RAIN_MULTI_LEVEL_ISSUED",
                "display_name": "호우특보 발효 안내",
                "button_label": "☔ 호우 복합",
                "enabled": False,
                "automation": False,
                "legacy": True,
                "hidden": True,
                "review_status": "confirmed",
                "review_label": "legacy — no exact Excel alias (HW-08 differs)",
                "aliases": [],
                "required_slots": ["지역", "기준시각"],
                "optional_slots": ["경보지역", "주의보지역", "홍수지역"],
                "text": (
                    "📢 {지역} 호우특보 발효 안내\n"
                    "({기준시각} 기준)\n"
                    "\n"
                    "{#경보지역}☔ 호우경보\n"
                    "{경보지역}\n"
                    "\n"
                    "{/경보지역}{#주의보지역}☔ 호우주의보\n"
                    "{주의보지역}\n"
                    "\n"
                    "{/주의보지역}{#홍수지역}☔ 홍수주의보\n"
                    "{홍수지역}\n"
                    "\n"
                    "{/홍수지역}하천 수위 상승과 저지대·지하공간 침수가 우려됩니다.\n"
                    "\n"
                    "하천변·산책로·둔치·공사장 등\n"
                    "위험지역 출입을 자제하고,\n"
                    "대피 권고 시 즉시 안전한 곳으로 이동해 주세요.\n"
                    "\n"
                    "📌 실시간 재난예보·행동요령\n"
                    "서울안전누리\n"
                    "👉 safecity.seoul.go.kr\n"
                ),
            }
        ],
    }


class _LiteralStr(str):
    """Force PyYAML to emit a literal block scalar."""


def _literal_representer(dumper: yaml.Dumper, data: _LiteralStr):
    return dumper.represent_scalar("tag:yaml.org,2002:str", str(data), style="|")


yaml.add_representer(_LiteralStr, _literal_representer)


def _prepare_for_dump(entry: dict) -> dict:
    out = dict(entry)
    text = out.get("text")
    if isinstance(text, str):
        # Ensure a trailing newline so `|` block form is stable.
        out["text"] = _LiteralStr(text if text.endswith("\n") else text + "\n")
    return out


def build_yaml_document(sync: SyncResult) -> dict:
    sections = _system_and_legacy_sections()
    return {
        "version": 2,
        "source": {
            "workbook": "templates/서울시_재난특보_X템플릿.xlsx",
            "catalog_sheet": CATALOG_SHEET,
            "wording_sheet": WORDING_SHEET,
        },
        "templates": [_prepare_for_dump(t) for t in sync.templates],
        "system_templates": [_prepare_for_dump(t) for t in sections["system_templates"]],
        "legacy_templates": [_prepare_for_dump(t) for t in sections["legacy_templates"]],
    }


def dump_yaml(document: dict) -> str:
    return yaml.dump(
        document,
        allow_unicode=True,
        default_flow_style=False,
        sort_keys=False,
        width=1000,
    )


def _yaml_normalized_compare(a: str, b: str) -> bool:
    """Structural equality ignoring incidental formatting differences."""
    return yaml.safe_load(a) == yaml.safe_load(b)


def validate_loader_accepts(yaml_text: str, tmp_path: Path) -> list[ValidationIssue]:
    issues: list[ValidationIssue] = []
    path = tmp_path / "message_templates.check.yaml"
    path.write_text(yaml_text, encoding="utf-8")
    clear_template_cache()
    try:
        loaded = load_templates(path)
    except Exception as exc:  # noqa: BLE001
        issues.append(
            ValidationIssue("", "yaml", 0, "loader", f"generated YAML failed to load: {exc}")
        )
        return issues
    finally:
        clear_template_cache()
    if "HW-01" not in loaded or "ORIGINAL_ONLY" not in loaded:
        issues.append(
            ValidationIssue(
                "", "yaml", 0, "loader", "generated YAML missing expected template ids"
            )
        )
    return issues


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--check", action="store_true", help="validate only; no writes")
    mode.add_argument("--write", action="store_true", help="write YAML when valid")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=TEMPLATES_PATH)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    workbook = args.workbook
    if not workbook.is_absolute():
        workbook = PROJECT_ROOT / workbook
    output = args.output
    if not output.is_absolute():
        output = PROJECT_ROOT / output

    sync = build_from_workbook(workbook)
    if not sync.ok:
        print("INVALID workbook / catalog:", file=sys.stderr)
        for issue in sync.issues:
            print(f"  - {issue.format()}", file=sys.stderr)
        return 1

    document = build_yaml_document(sync)
    yaml_text = dump_yaml(document)

    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        loader_issues = validate_loader_accepts(yaml_text, Path(tmp))
    if loader_issues:
        print("INVALID generated YAML:", file=sys.stderr)
        for issue in loader_issues:
            print(f"  - {issue.format()}", file=sys.stderr)
        return 1

    print(
        f"workbook OK: total={sync.total} automation={sync.automation_count} "
        f"reference={sync.reference_count}"
    )

    current = output.read_text(encoding="utf-8") if output.exists() else ""
    in_sync = _yaml_normalized_compare(current, yaml_text) if current else False

    if args.check:
        if in_sync:
            print(f"YAML synchronized with workbook: {output}")
            return 0
        print(f"YAML STALE relative to workbook: {output}", file=sys.stderr)
        return 1

    # --write
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(yaml_text, encoding="utf-8")
    clear_template_cache()
    print(f"wrote {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
